from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Gorilla_Mutantz_Gym_Client_Progress_Report_2026-06-20.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(20, 24, 33)
MUTED = RGBColor(96, 103, 118)
LIGHT_FILL = "F2F4F7"
BLUE_FILL = "E8EEF5"
WHITE = "FFFFFF"
BORDER = "D9DEE7"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = BORDER, size: str = "6") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = margins.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            margins.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            set_cell_border(cell)


def set_run(run, size=11, bold=None, color=None, italic=None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_para(doc, text="", style=None, size=11, bold=False, color=INK, italic=False, after=6, before=0, align=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        set_run(run, size=size, bold=bold, color=color, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_bullets(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        run = p.add_run(item)
        set_run(run, size=11, color=INK)


def add_numbered(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        run = p.add_run(item)
        set_run(run, size=11, color=INK)


def add_label_table(doc, rows: list[tuple[str, str]], widths=(1.75, 4.75)) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_width(table, list(widths))
    for idx, (label, value) in enumerate(rows):
        cells = table.rows[idx].cells
        set_cell_shading(cells[0], LIGHT_FILL)
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p0 = cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        set_run(r0, size=10, bold=True, color=DARK_BLUE)
        p1 = cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(value)
        set_run(r1, size=10, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_matrix(doc, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_width(table, widths)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, BLUE_FILL)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run(r, size=9.5, bold=True, color=DARK_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            set_cell_shading(cell, WHITE)
            set_cell_border(cell)
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            set_run(r, size=9.2, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def add_callout(doc, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [6.5])
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "F4F6F9")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    set_run(r, size=10.5, bold=True, color=DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    set_run(r2, size=10, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.167

    header = section.header.paragraphs[0]
    header.text = "Gorilla Mutantz Gym Access System | Progress Report"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(header.runs[0], size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.text = "Prepared for client review | v0.1 progress snapshot"
    set_run(footer.runs[0], size=9, color=MUTED)

    return doc


def add_title_page(doc: Document) -> None:
    add_para(doc, "PROJECT PROGRESS REPORT", size=10, bold=True, color=BLUE, after=2)
    add_para(doc, "Gorilla Mutantz Gym Access System", size=24, bold=True, color=INK, after=4)
    add_para(doc, "Client progress summary for Login, Dashboard, Members, and Membership modules", size=13, color=MUTED, after=18)
    add_label_table(doc, [
        ("Prepared for", "Client review and progress update"),
        ("Prepared on", "20 June 2026"),
        ("System version", "v0.1 progress build"),
        ("Design direction", "Metronic-inspired light administration interface"),
        ("Deployment context", "Local gym environment through Herd, with PostgreSQL backend and local HTTPS support"),
    ])
    add_callout(
        doc,
        "Overall progress",
        "The project foundation, authentication, dashboard reporting surface, member registration/editing, member worklists, profile view, and membership package/assignment flows are implemented at working prototype level. The current build is designed for front-desk operation and is already enforcing role permissions, audit logging, phone validation, password security, membership expiry rules, and theme persistence.",
    )
    add_heading(doc, "Modules covered in this report", 2)
    add_numbered(doc, [
        "Login Page and access security",
        "Dashboard / Daily Sales Report",
        "Members Page and member registration workflow",
        "Membership Page, packages, assignment, renewal, and expiry worklists",
    ])
    doc.add_page_break()


def section_login(doc: Document) -> None:
    add_heading(doc, "1. Login Page", 1)
    add_para(doc, "The login page is complete as the secure entry point for staff and administrators. It uses the Gorilla Mutantz Gym brand, a focused sign-in panel, and a support modal for password reset requests.")
    add_heading(doc, "User-facing layout", 2)
    add_bullets(doc, [
        "Brand block: GM mark, Gorilla Mutantz Gym name, and Gym access system subtitle.",
        "Security label: Secure access.",
        "Primary heading: Sign in to dashboard.",
        "Supporting copy: Manage members, cards, receipts, and access control from one workspace.",
        "Right-side system status panel: Access Online, 2 door access units, active access count, last sync time, and pending jobs.",
        "Dark/light theme is detected before page render to avoid flash and to respect saved user preference.",
    ])
    add_heading(doc, "Login form fields and controls", 2)
    add_matrix(doc, ["Field / Control", "Purpose", "Validation / Behavior"], [
        ["Username", "Identifies staff/admin account.", "Required, string, max 100 characters, normalized to lowercase for authentication."],
        ["Password", "Authenticates the user.", "Required, max 255 characters, protected as password input."],
        ["Remember me", "Keeps the session active according to Laravel remember-token behavior.", "Optional checkbox."],
        ["Forgot password?", "Opens support modal instead of automated reset.", "Modal instructs staff to contact support via WhatsApp."],
        ["Sign In", "Submits credentials.", "Regenerates session after successful login to reduce session fixation risk."],
    ], [1.4, 2.35, 2.75])
    add_heading(doc, "Password reset support modal", 2)
    add_bullets(doc, [
        "Modal title: Contact support to reset password.",
        "Reset process is intentionally manual for security.",
        "Support contact is WhatsApp only: +601128520309.",
        "WhatsApp link uses https://wa.me/601128520309.",
        "Staff are asked to include staff name and branch when requesting reset support.",
    ])
    add_heading(doc, "Security behavior implemented", 2)
    add_bullets(doc, [
        "Only active users can sign in.",
        "Failed login attempts are rate-limited by username and IP address.",
        "Too many attempts trigger a temporary lockout message.",
        "Credentials use username and password, not email login.",
        "Successful login records last_login_at.",
        "Successful login records an audit trail entry for auth/login.",
        "Logout records auth/logout, invalidates the session, and regenerates CSRF token.",
        "Theme preference survives logout/login and is stored per user when authenticated.",
        "Authenticated routes require active user middleware.",
        "Feature pages are protected by permission middleware such as members.manage and memberships.manage.",
    ])


def section_dashboard(doc: Document) -> None:
    add_heading(doc, "2. Dashboard / Daily Sales Report", 1)
    add_para(doc, "The current dashboard is implemented as a Daily Sales Report surface, following the provided sample dashboard content while retaining the Metronic-inspired visual system. It is designed to help the front desk and owner review sales performance, payment collection, and transaction details at a glance.")
    add_heading(doc, "Dashboard controls", 2)
    add_matrix(doc, ["Control", "Current behavior", "Client value"], [
        ["Date", "Read-only date field in the report filter panel.", "Defines the daily sales period shown in the dashboard."],
        ["Outlet", "Hidden while multi-branch mode is not active.", "Will become visible later when multi-branch support is enabled."],
        ["Generate Report", "Primary action button for report refresh.", "Supports manual report refresh workflow."],
        ["Export", "Icon action with hover guide.", "Prepared for report export workflow."],
        ["Update deployment", "Admin-only icon action guarded by deployment key, rate limit, and executable script checks.", "Allows controlled local deployment update when configured."],
    ], [1.45, 2.55, 2.5])
    add_heading(doc, "Dashboard KPI cards", 2)
    add_matrix(doc, ["Measurement", "Displayed / calculated value", "Meaning"], [
        ["Total Revenue", "Daily completed sales total.", "Total money collected from all completed sales for the selected date."],
        ["Membership Sales", "Daily completed membership sale and renewal total.", "Shows how much revenue came from membership activity."],
        ["Product Sales", "Daily completed product sale total.", "Shows merchandise/product revenue."],
        ["Cash Collection", "Daily cash payment total.", "Shows physical cash collected."],
        ["Online Payment", "Daily online/QR/bank transfer payment total.", "Shows non-cash collection total."],
        ["Monthly Revenue", "Completed sales between current month start and month end.", "Backend metric available for management view."],
    ], [1.55, 2.25, 2.7])
    add_heading(doc, "Dashboard charts and breakdowns", 2)
    add_bullets(doc, [
        "Sales Breakdown donut: separates Membership Sales, Product Sales, and Other Sales by amount and percentage.",
        "Payment Method Breakdown donut: separates Cash and Online Payment values by amount and percentage.",
        "Weekly Sales Trend line chart: displays seven-day sales movement from Monday to Sunday.",
        "The weekly backend metric separates membership sales and product sales for each of the last seven days.",
    ])
    add_heading(doc, "Sales Details table", 2)
    add_matrix(doc, ["Column", "Captured / displayed data"], [
        ["Time", "Transaction time, e.g. 09:15 AM."],
        ["Receipt No.", "Receipt number such as INV-2026-0611-0001."],
        ["Type", "Sale classification such as Membership Sale or Product Sale."],
        ["Description", "Human-readable transaction description, including package/member or product quantity."],
        ["Category", "Transaction category, e.g. Membership, Supplements, Drinks, Snacks."],
        ["Payment Method", "Cash, Online Payment, QR, bank transfer, or other enabled method."],
        ["Amount (RM)", "Transaction amount in Ringgit Malaysia."],
        ["Received By", "Staff/admin user who received or processed the transaction."],
    ], [1.55, 4.95])
    add_heading(doc, "Backend dashboard measurements available", 2)
    add_matrix(doc, ["Metric group", "Measurements captured"], [
        ["Member KPIs", "Total members, active members, expired members, and expiring soon members."],
        ["Sales KPIs", "Today sales, today membership sales, today product sales, and monthly revenue."],
        ["Membership status overview", "Active, Expiring Soon, Expired, Suspended, and Cancelled membership counts."],
        ["Weekly sales chart", "Last seven days, split into membership sales and product sales."],
        ["Recent activity", "Latest 10 audit log entries, including module, action, user, and timestamp."],
        ["System status", "Database online/offline, backup path status, 2-unit door access configuration, last access sync time/status, and pending access sync jobs."],
    ], [1.9, 4.6])
    add_callout(doc, "Important current dashboard note", "The visible dashboard is already styled and structured as the daily report requested. Some values in the visual sample are currently static demo values while backend metric services are implemented for live calculation as the remaining POS/reporting flows are connected.")


def section_members(doc: Document) -> None:
    add_heading(doc, "3. Members Page", 1)
    add_para(doc, "The Members module is implemented as the main front-desk workflow for registering, searching, editing, reviewing, and maintaining gym member records. It includes member registration, worklists for expiring/expired memberships, photo capture, profile view, and controlled suspend/reactivate actions.")
    add_heading(doc, "Members sidebar and pages", 2)
    add_bullets(doc, [
        "Manage Members: searchable member list with row actions.",
        "Member Registration: full add-new-member form with membership setup and photo/RFID sections.",
        "Expiring Soon: members whose active membership ends within the configured expiring_soon_days window.",
        "Expired Members: members whose membership is expired or whose active membership end date has passed.",
        "Photo Capture: members without a stored photo.",
    ])
    add_heading(doc, "Manage Members table columns", 2)
    add_matrix(doc, ["Column", "Description"], [
        ["Member No.", "Auto-generated Gorilla Mutantz Gym member number, using GM + year/month + sequence."],
        ["Name", "Member full name."],
        ["Phone", "Primary phone number."],
        ["IC / Passport", "Identity document number where provided."],
        ["Status", "Displayed status: Active, Expiring, Suspended, or Inactive. Expiring is calculated from membership date rules."],
        ["Updated", "Last updated date."],
        ["Actions", "Icon-only actions with hover guide: view profile and edit member."],
    ], [1.55, 4.95])
    add_heading(doc, "Search and filtering", 2)
    add_bullets(doc, [
        "Search supports member number, full name, phone, and IC/passport number.",
        "Worklists reuse the member list layout for consistent operation.",
        "Expiring Soon uses membership end date from today through today plus expiring_soon_days.",
        "Expired Members includes memberships with explicit expired status or active memberships with end date before today.",
        "Photo Capture filters members where photo_path is empty.",
    ])
    add_heading(doc, "Member registration form sections and fields", 2)
    add_matrix(doc, ["Section", "Fields captured"], [
        ["Personal Information", "Full Name, IC / Passport No., Date of Birth, Gender, Phone Number, Email, Address."],
        ["Emergency Contact", "Contact Name, Relationship, Contact Number."],
        ["Additional Information", "Join Date auto-generated on save, Referred By active-member filter, Notes placeholder."],
        ["Membership Information", "Membership Type, Start Date, End Date, Amount (RM), Payment Method, Payment Status."],
        ["Photo & RFID Card", "Member Photo from camera capture or upload, RFID Card Number manually entered."],
        ["Remarks", "Internal remarks for staff/admin use."],
    ], [1.8, 4.7])
    add_heading(doc, "Member registration validation and rules", 2)
    add_bullets(doc, [
        "Full Name is required, max 255 characters.",
        "Phone Number is required and must match Malaysian phone formats such as +601128520309, 011-2852 0309, or +6065252503.",
        "Emergency Contact Number uses the same Malaysian phone validator when provided.",
        "IC / Passport No. must be unique when provided.",
        "Gender supports Male and Female.",
        "Email must be valid email format when provided.",
        "Address and Remarks allow up to 2,000 characters.",
        "RFID Card Number must be unique when provided.",
        "Referrer must be an active member and cannot be the same member during edit.",
        "Photo upload must be an image and max 2 MB.",
        "Captured camera image is stored from base64 image data.",
    ])
    add_heading(doc, "Photo capture and upload behavior", 2)
    add_bullets(doc, [
        "Camera can start from the member form when browser permissions allow it.",
        "After Capture, the camera automatically stops.",
        "Captured photo is immediately shown in the member photo preview frame.",
        "Uploaded photo is also shown immediately in the same preview frame.",
        "When editing an existing member, the stored photo appears in the preview frame automatically.",
        "The system supports local HTTPS so camera permissions can work in the local gym environment.",
    ])
    add_heading(doc, "RFID handling in current scope", 2)
    add_bullets(doc, [
        "RFID scan button/automatic card scanning is intentionally removed for now.",
        "Staff manually key in the card number at the counter.",
        "Card Status display is hidden for now but underlying card/access logic is preserved for future RFID module activation.",
        "Manual entry is more practical because the Dahua device is at the door, not the counter.",
    ])
    add_heading(doc, "Member edit behavior", 2)
    add_bullets(doc, [
        "Edit member page shows existing personal, contact, photo, RFID, referral, and remarks data.",
        "Membership Information shows the latest existing membership package, start date, end date, amount, payment method display, and payment status.",
        "End Date is editable on member edit to allow correction of expiry date.",
        "Saving an updated expiry date updates the latest membership record and writes an audit log action: expiry_updated.",
    ])
    add_heading(doc, "Member profile view", 2)
    add_bullets(doc, [
        "Profile header shows member photo/initials, full name, member number, phone, email, and current display status.",
        "Display status can show Expiring when latest active membership ends within the expiring-soon window.",
        "Summary cards show Membership, Expiry Date, RFID Card, and Access Sync.",
        "Detailed sections show Contact, Personal Information, Membership & Access, Emergency Contact, Referral Source, and Remarks.",
        "Profile actions include edit profile, assign/renew membership, suspend/reactivate member, and suspend membership where permitted.",
    ])
    add_heading(doc, "Member status labels", 2)
    add_matrix(doc, ["Displayed status", "Meaning"], [
        ["Active", "Member record is active and not inside the expiring-soon membership window."],
        ["Expiring", "Member is active and latest active membership ends within expiring_soon_days."],
        ["Suspended", "Member has been manually suspended."],
        ["Inactive", "Member record is disabled/non-active."],
    ], [1.7, 4.8])


def section_membership(doc: Document) -> None:
    add_heading(doc, "4. Membership Page", 1)
    add_para(doc, "The Membership module manages membership packages and member membership lifecycle actions. It is already protected by membership permissions and records audit/access-sync activity when memberships are assigned, renewed, or suspended.")
    add_heading(doc, "Membership package list", 2)
    add_matrix(doc, ["Column", "Description"], [
        ["Name", "Package name, e.g. Monthly, Quarterly, Yearly, Walk-in."],
        ["Duration", "Duration in days."],
        ["Price", "Package price in RM."],
        ["Access", "Whether the package allows door access."],
        ["Status", "Active or Inactive."],
        ["Actions", "Icon-only edit action with hover guide."],
    ], [1.45, 5.05])
    add_heading(doc, "Membership package form fields", 2)
    add_matrix(doc, ["Field", "Purpose", "Validation"], [
        ["Package Name", "Defines the display name for staff.", "Required, unique, max 120 characters."],
        ["Duration Days", "Defines membership length.", "Required integer, min 1, max 3650."],
        ["Price (RM)", "Defines default package selling price.", "Required numeric, min 0, max 999999.99."],
        ["Status", "Controls whether package is selectable.", "Required: Active or Inactive."],
        ["Walk-in package", "Flags package as walk-in style.", "Optional boolean."],
        ["Allow door access", "Controls whether membership queues enable-card access sync.", "Optional boolean, default enabled."],
    ], [1.55, 2.25, 2.7])
    add_heading(doc, "Assign / renew membership form fields", 2)
    add_matrix(doc, ["Field", "Purpose"], [
        ["Package", "Selects active membership package and shows price/duration."],
        ["Start Date / Renewal Date", "Sets when the membership period starts or renewal is processed."],
        ["Amount (RM)", "Allows staff to record the charged membership amount."],
        ["Payment Status", "Paid or Unpaid."],
        ["Current expiry note", "Shown during renewal; explains renewal extension from current expiry when renewal is before expiry."],
    ], [1.75, 4.75])
    add_heading(doc, "Membership expiry calculation", 2)
    add_callout(doc, "Implemented expiry rule", "The expiry date represents the last active day of the membership. A 30-day membership starting 19 Jun 2026 expires on 18 Jul 2026. This rule is applied to member registration, membership assignment, membership renewal, and POS membership sale/renewal.")
    add_bullets(doc, [
        "New membership: end date = start date + duration days - 1.",
        "Renewal before existing expiry: new end date extends from the existing expiry date.",
        "Renewal after expiry: new end date starts from the new renewal/start date using the same inclusive rule.",
        "The edit member form allows expiry date adjustment for correction and audits the change.",
    ])
    add_heading(doc, "Membership worklist pages", 2)
    add_matrix(doc, ["Page", "Logic", "Columns"], [
        ["Expiring Soon", "Active memberships with end_date from today through today + expiring_soon_days.", "Member, Package, Start, Expiry, Status, Profile action."],
        ["Expired Members", "Memberships marked expired, or active memberships with end_date before today.", "Member, Package, Start, Expiry, Status, Profile action."],
    ], [1.4, 2.7, 2.4])
    add_heading(doc, "Membership statuses", 2)
    add_matrix(doc, ["Status", "Meaning"], [
        ["Active", "Membership is currently valid."],
        ["Expiring Soon", "Calculated management bucket for active memberships near expiry."],
        ["Expired", "Membership end date has passed or status is explicitly expired."],
        ["Suspended", "Membership manually suspended."],
        ["Cancelled", "Membership cancelled before normal expiry."],
    ], [1.6, 4.9])
    add_heading(doc, "Audit and access sync behavior", 2)
    add_bullets(doc, [
        "Package create/update records audit actions package_created and package_updated.",
        "Membership assignment records assigned audit action.",
        "Membership renewal records renewed audit action.",
        "Membership suspension records suspended audit action.",
        "Manual expiry correction from member edit records expiry_updated audit action.",
        "Assignment/renewal creates pending access sync logs for both door access units.",
        "Access sync action is ENABLE_CARD when membership is active and package allows access; otherwise DISABLE_CARD. The sync is considered successful only after both units accept the update.",
    ])


def add_closing(doc: Document) -> None:
    add_heading(doc, "Current Client-Visible Completion Summary", 1)
    add_matrix(doc, ["Area", "Progress status", "Client-visible result"], [
        ["Login Page", "Implemented", "Secure username/password login, remember me, manual password support modal, theme persistence, audit logging."],
        ["Dashboard", "Implemented visual/report surface; backend metric service prepared", "Daily sales report layout with KPI cards, charts, transaction table, admin deployment update control."],
        ["Members Page", "Implemented", "Searchable members list, registration/edit forms, photo capture/upload preview, RFID manual entry, profile page, worklists."],
        ["Membership Page", "Implemented core lifecycle", "Package CRUD, assign membership, renew membership, expiry rule, expiring/expired worklists, audit/access sync."],
    ], [1.45, 1.75, 3.3])
    add_heading(doc, "Recommended next client demo flow", 2)
    add_bullets(doc, [
        "Show login and explain manual password reset via WhatsApp support.",
        "Open dashboard and walk through KPIs, daily sales charting, payment breakdown, and transaction table.",
        "Register a member, including phone validation, active referrer search, photo capture/upload, RFID manual entry, and initial membership setup.",
        "Edit the member and demonstrate existing photo preview plus editable membership end date.",
        "Open the profile page and explain Active/Expiring/Suspended statuses.",
        "Show membership package list, create/edit package form, assign/renew membership, and expiry rule.",
    ])


def build() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = setup_document()
    add_title_page(doc)
    section_login(doc)
    section_dashboard(doc)
    section_members(doc)
    section_membership(doc)
    add_closing(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
